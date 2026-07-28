"""匯出:summary 報表(工具計算結果)與 v2 範本(多機型格式)。"""
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from captool.models import PlanInput
from captool.planner import PlanResult

_ARIAL = Font(name="Arial")
_BOLD = Font(name="Arial", bold=True)
_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

_SUMMARY_HEADERS = ["Fab", "BM Group", "月份", "需求 vcore", "VM 顆數", "進機",
                    "退回", "新啟用機台", "建議採購", "月末總機台", "月末空機",
                    "剩餘可售 vcore", "狀態", "缺口 vcore"]


def _fmt(d: dict[str, int]) -> str:
    return "; ".join(f"{k}×{v}" for k, v in sorted(d.items())) if d else ""


def _style_row(ws, row_idx: int, fill=None) -> None:
    for cell in ws[row_idx]:
        cell.font = _ARIAL
        if fill is not None:
            cell.fill = fill


def export_summary(plan_input: PlanInput, plan_result: PlanResult, path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(_SUMMARY_HEADERS)
    _style_row(ws, 1)
    for cell in ws[1]:
        cell.font = _BOLD
    for o in plan_result.outcomes:
        ws.append([
            o.pool.fab, o.pool.bm_group, o.month, o.demand_vcore,
            sum(c for _, c in o.vm_demand),
            _fmt(o.movein), _fmt(o.returns), _fmt(o.machines_opened),
            _fmt(o.suggested_purchases), _fmt(o.stock_total), _fmt(o.stock_empty),
            round(o.free_vcore_total, 2),
            "OK" if o.feasible else "缺口",
            round(o.shortfall_vcore, 2) if not o.feasible else 0,
        ])
        _style_row(ws, ws.max_row, fill=None if o.feasible else _RED)
    ws.append([])
    ws.append([f"本表由容量規劃工具產出(模式:{plan_result.mode}),數值為 solver 計算結果,非公式。"])
    _style_row(ws, ws.max_row)

    wp = wb.create_sheet("placements")
    wp.append(["Fab", "BM Group", "月份", "機型", "新啟用台數"])
    _style_row(wp, 1)
    for cell in wp[1]:
        cell.font = _BOLD
    for o in plan_result.outcomes:
        for sku_name, cnt in sorted(o.machines_opened.items()):
            wp.append([o.pool.fab, o.pool.bm_group, o.month, sku_name, cnt])
            _style_row(wp, wp.max_row)
    wb.save(path)


def generate_v2_template(plan_input: PlanInput, path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "README"
    lines = [
        "容量規劃 v2 範本填表說明(新版面)",
        "",
        "各廠區 tab —— 需求區(A 欄起):",
        "  A=Product, B=BM Group, C=VM vcore, D=爆炸半徑(1:x), E=Tenant, F 欄起為各月 vcore。",
        "  C(VM vcore)空 = 粗粒度需求(整包 vcore);有值 = 每顆 VM 的 vcore 尺寸。",
        "  D(爆炸半徑 1:x)= 一台實體機最多住幾顆該 VM(EX: 2);空=不限。",
        "  E(Tenant)= 自由(留空,可與其他自由 product 混) / 指定群組名(同名才可共用) / 獨佔。",
        "  每月一律填 vcore;detail product 的顆數由 vcore ÷ VM vcore 反推(非整數倍會進位並提示)。",
        "  同 product 若配置改變,請拆成兩列(如 A-1:1、A-1:2)。",
        "退回區(T 欄起):T=Product, U=BM Group, V=機型, W 欄起為各月退回台數。",
        "HW_SKU / HW_Current / HW_MoveIn:機型目錄 / 期初庫存 / 進機計畫(硬體 team 維護)。",
        "注意:機型名稱須存在於 HW_SKU;月份一律 YYYY-MM;示範列(product 含『示範』)請刪除。",
    ]
    for line in lines:
        ws.append([line])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = _ARIAL
    ws["A1"].font = _BOLD

    fabs = sorted({p.fab for p in plan_input.pools})
    for fab in fabs:
        wf = wb.create_sheet(fab)
        wf["A4"], wf["B4"] = "Product", "BM Group"
        wf["C4"] = "VM vcore"
        wf["D4"] = "爆炸半徑 (1:x), EX: 2"
        wf["E4"] = "Tenant (自由=留空/指定群組/獨佔)"
        wf.cell(row=4, column=20, value="Product")
        wf.cell(row=4, column=21, value="BM Group")
        wf.cell(row=4, column=22, value="機型")
        for i, month in enumerate(plan_input.months):
            wf.cell(row=4, column=6 + i, value=month)     # 需求月份 F 起
            wf.cell(row=4, column=23 + i, value=month)    # 退回月份 W 起

        # 帶入既有需求(粗粒度)
        fab_demands = [d for d in plan_input.demands if d.pool.fab == fab]
        rows: dict[tuple[str, str], int] = {}
        r = 5
        for d in fab_demands:
            key = (d.product, d.pool.bm_group)
            if key not in rows:
                rows[key] = r
                wf.cell(row=r, column=1, value=d.product)
                wf.cell(row=r, column=2, value=d.pool.bm_group)
                r += 1
            wf.cell(row=rows[key], column=6 + plan_input.months.index(d.month),
                    value=d.vcore)

        # 帶入既有 VM-detail policy + 對應需求(每月顆數 × VM vcore)
        fab_policies = [p for p in plan_input.policies if p.pool.fab == fab]
        vm_counts: dict[tuple, int] = defaultdict(int)
        for v in plan_input.vm_demands:
            if v.pool.fab == fab:
                vm_counts[(v.pool, v.product, v.month)] += v.count
        for pol in fab_policies:
            wf.cell(row=r, column=1, value=pol.product)
            wf.cell(row=r, column=2, value=pol.pool.bm_group)
            wf.cell(row=r, column=3, value=pol.vm_size_vcore)
            if pol.max_per_machine is not None:
                wf.cell(row=r, column=4, value=pol.max_per_machine)
            if pol.co_residency == "exclusive":
                wf.cell(row=r, column=5, value="獨佔")
            elif pol.co_residency != "free":
                wf.cell(row=r, column=5, value=pol.co_residency)
            for i, month in enumerate(plan_input.months):
                count = vm_counts.get((pol.pool, pol.product, month), 0)
                if count and pol.vm_size_vcore is not None:
                    wf.cell(row=r, column=6 + i, value=count * pol.vm_size_vcore)
            r += 1

        # 一列示範 detail
        wf.cell(row=r, column=1, value="示範product(可刪)")
        wf.cell(row=r, column=2, value="network1")
        wf.cell(row=r, column=3, value=60)
        wf.cell(row=r, column=4, value=1)
        wf.cell(row=r, column=5, value="teamA")
        wf.cell(row=r, column=6, value=120)
        for c in range(1, 7):
            wf.cell(row=r, column=c).fill = _YELLOW

        # 帶入既有 return(機型填 default)
        fab_returns = [x for x in plan_input.returns if x.pool.fab == fab]
        rrows: dict[tuple[str, str, str], int] = {}
        r = 5
        for x in fab_returns:
            key = (x.product, x.pool.bm_group, x.sku_name)
            if key not in rrows:
                rrows[key] = r
                wf.cell(row=r, column=20, value=x.product)
                wf.cell(row=r, column=21, value=x.pool.bm_group)
                wf.cell(row=r, column=22, value=x.sku_name)
                r += 1
            wf.cell(row=rrows[key], column=23 + plan_input.months.index(x.month),
                    value=x.count)

        for row in wf.iter_rows():
            for cell in row:
                if cell.font is not _BOLD:
                    cell.font = _ARIAL
        for c in range(1, 6):
            wf.cell(row=4, column=c).font = _BOLD
        for c in (20, 21, 22):
            wf.cell(row=4, column=c).font = _BOLD

    ws1 = wb.create_sheet("HW_SKU")
    ws1.append(["name", "vcore_per_node", "usable_ratio"])
    for sku in plan_input.skus.values():
        ws1.append([sku.name, sku.vcore_per_node, sku.usable_ratio])
    ws2 = wb.create_sheet("HW_Current")
    ws2.append(["fab", "bm_group", "sku", "count"])
    for c in plan_input.currents:
        ws2.append([c.pool.fab, c.pool.bm_group, c.sku_name, c.count])
    ws3 = wb.create_sheet("HW_MoveIn")
    ws3.append(["fab", "bm_group", "sku", "month", "count"])
    for m in plan_input.moveins:
        ws3.append([m.pool.fab, m.pool.bm_group, m.sku_name, m.month, m.count])
    for sheet in (ws1, ws2, ws3):
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = _ARIAL
        for cell in sheet[1]:
            cell.font = _BOLD
    wb.save(path)
