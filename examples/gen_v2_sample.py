"""產生新版面 v2 sample(含假資料),並驗證可匯入 + 推演。

用法:.venv\\Scripts\\python.exe examples\\gen_v2_sample.py
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

BOLD = Font(bold=True)
YELLOW = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
MONTHS = ["2026-07", "2026-08", "2026-09"]

SKUS = [("std-64", 64, 0.8), ("big-128", 128, 0.8), ("gpu-80", 80, 0.8)]
# 需求列:(fab, product, group, vm_size|None, max_per|None, coresid, {month: vcore})
DEMAND_ROWS = [
    ("A", "web-svc", "network1", None, None, "", {"2026-07": 200, "2026-08": 120}),           # 粗粒度
    ("A", "AI-train", "network1", 96, 1, "獨佔", {"2026-07": 192}),                            # detail 獨佔,2 顆
    ("A", "cache", "network1", 60, 2, "teamA", {"2026-07": 180, "2026-08": 120}),              # detail 群組
    ("B", "db", "network1", 48, 2, "", {"2026-07": 144, "2026-08": 96, "2026-09": 48}),        # detail 自由
]
CURRENT = [("A", "network1", "std-64", 6), ("A", "network1", "big-128", 3),
           ("B", "network1", "std-64", 8)]
MOVEIN = [("B", "network1", "std-64", "2026-08", 2)]
RETURNS = [("B", "network1", "db", "std-64", "2026-09", 1)]


def build(path):
    wb = Workbook(); wb.remove(wb.active)
    fabs = sorted({r[0] for r in DEMAND_ROWS} | {c[0] for c in CURRENT})
    for fab in fabs:
        wf = wb.create_sheet(fab)
        wf["A4"], wf["B4"] = "Product", "BM Group"
        wf["C4"] = "VM vcore"
        wf["D4"] = "爆炸半徑 (1:x), EX: 2"
        wf["E4"] = "Tenant (自由=留空/指定群組/獨佔)"
        wf.cell(row=4, column=20, value="Product")
        wf.cell(row=4, column=21, value="BM Group")
        wf.cell(row=4, column=22, value="機型")
        for i, m in enumerate(MONTHS):
            wf.cell(row=4, column=6 + i, value=m)
            wf.cell(row=4, column=23 + i, value=m)
        r = 5
        for (f, prod, grp, size, mx, cor, mvals) in DEMAND_ROWS:
            if f != fab:
                continue
            wf.cell(row=r, column=1, value=prod)
            wf.cell(row=r, column=2, value=grp)
            if size is not None:
                wf.cell(row=r, column=3, value=size)
                if mx is not None:
                    wf.cell(row=r, column=4, value=mx)
                if cor:
                    wf.cell(row=r, column=5, value=cor)
                for c in range(1, 6):
                    wf.cell(row=r, column=c).fill = YELLOW
            for i, m in enumerate(MONTHS):
                if m in mvals:
                    wf.cell(row=r, column=6 + i, value=mvals[m])
            r += 1
        r = 5
        for (f, grp, prod, sku, month, cnt) in RETURNS:
            if f != fab:
                continue
            wf.cell(row=r, column=20, value=prod)
            wf.cell(row=r, column=21, value=grp)
            wf.cell(row=r, column=22, value=sku)
            wf.cell(row=r, column=23 + MONTHS.index(month), value=cnt)
            r += 1
        for c in range(1, 6):
            wf.cell(row=4, column=c).font = BOLD
        for c in (20, 21, 22):
            wf.cell(row=4, column=c).font = BOLD

    hs = wb.create_sheet("HW_SKU"); hs.append(["name", "vcore_per_node", "usable_ratio"])
    for s in SKUS:
        hs.append(list(s))
    hc = wb.create_sheet("HW_Current"); hc.append(["fab", "bm_group", "sku", "count"])
    for row in CURRENT:
        hc.append(list(row))
    hm = wb.create_sheet("HW_MoveIn"); hm.append(["fab", "bm_group", "sku", "month", "count"])
    for row in MOVEIN:
        hm.append(list(row))
    for sheet in (hs, hc, hm):
        for cell in sheet[1]:
            cell.font = BOLD
    wb.save(path)


def verify(path):
    from captool.importer import import_v2
    from captool.planner import run_check
    from captool.solver.naive import NaiveSolver
    pi = import_v2(path)
    print("機型:", list(pi.skus), "pools:", [p.label for p in pi.pools], "月份:", pi.months)
    print("政策:")
    for pol in pi.policies:
        print(f"  {pol.pool.label} {pol.product}: VM {pol.vm_size_vcore}v "
              f"每台上限={pol.max_per_machine} 共居={pol.co_residency}")
    print("警告:", [i.message for i in pi.issues] or "無")
    res = run_check(pi, NaiveSolver(), mode="conservative")
    for o in res.outcomes:
        print(f"  [{'OK  ' if o.feasible else '缺口'}] {o.pool.label} {o.month}: "
              f"液體={o.demand_vcore} VM={o.vm_demand} 新啟用={o.machines_opened}")


if __name__ == "__main__":
    out = ROOT / "examples" / "v2_sample_new_format.xlsx"
    build(out)
    print("wrote", out, "\n")
    verify(out)
