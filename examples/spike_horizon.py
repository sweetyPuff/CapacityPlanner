# -*- coding: utf-8 -*-
"""Spike:把 v2_sample_solver.xlsx 的 worker 資料組成 CapacityPlanRequest,
POST 到本地 solver /v1/capacity/plan,印出 CapacityReport 重點。

worker-only(control-plane 菜單待 §6.2 role-enum/co-tenancy 敲定後再接)。
單維 vcore:BM cpu_cores = 機型 sellable(vcore_per_node×usable_ratio,floor),
mem/storage 設大值不綁;worker VM 規格 8 core、mem/storage=0,讓 cpu 為唯一約束。
每個 fab 一個 request(single-fab 模式 fab=""),network=bm_group。
"""
import json
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

SAMPLE = Path(__file__).resolve().parent / "v2_sample_solver.xlsx"
URL = "http://127.0.0.1:50051/v1/capacity/plan"
WORKER_VM = {"cpu_cores": 8, "memory_mib": 0, "storage_gb": 0}
BIG_MEM, BIG_STO = 10 ** 8, 10 ** 6
NON_FAB = {"HW_SKU", "HW_Current", "HW_MoveIn", "HW_Caps",
           "Cluster_Menu", "New_Build", "README", "summary"}

wb = load_workbook(SAMPLE, data_only=True)

# 機型 → sellable cpu
sku_cpu = {}
for r in wb["HW_SKU"].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    sku_cpu[str(r[0])] = int(int(r[1]) * float(r[2]))

# HW_Current: fab, bm_group, sku, count, ag
current = [(str(r[0]), str(r[1]), str(r[2]), int(r[3]), str(r[4]))
           for r in wb["HW_Current"].iter_rows(min_row=2, values_only=True)
           if r[0] is not None]

# HW_Caps: fab, network, ag, max_bm
caps = [(str(r[0]), str(r[1]), str(r[2]), int(r[3]))
        for r in wb["HW_Caps"].iter_rows(min_row=2, values_only=True)
        if r[0] is not None]

# worker 需求(粗粒度):fab tabs,月份 row4 自 F(6),product col1, group col2
demand = []   # (fab, network, product, month, vcore)
fab_sheets = [n for n in wb.sheetnames if n not in NON_FAB]
for name in fab_sheets:
    ws = wb[name]
    months = []
    c = 6
    while ws.cell(row=4, column=c).value is not None:
        months.append((c, str(ws.cell(row=4, column=c).value)))
        c += 1
    row = 5
    while ws.cell(row=row, column=1).value is not None:
        product = str(ws.cell(row=row, column=1).value)
        group = str(ws.cell(row=row, column=2).value)
        vm_size = ws.cell(row=row, column=3).value   # 空=worker 粗粒度
        for col, month in months:
            v = ws.cell(row=row, column=col).value
            if v and (vm_size in (None, "")):
                demand.append((name, group, product, month, int(v)))
        row += 1

fabs = sorted({d[0] for d in demand} | {c[0] for c in current})


def build_request(fab):
    db = [{
        "cluster_id": product, "node_role": "worker", "period": month,
        "cpu_cores": vcore, "network": network, "vm_specs": [WORKER_VM],
    } for (f, network, product, month, vcore) in demand if f == fab]

    in_stock = []
    for (f, network, sku, count, ag) in current:
        if f != fab:
            continue
        cpu = sku_cpu[sku]
        for k in range(count):
            in_stock.append({
                "id": f"{fab}-{network}-{sku}-{ag}-{k}",
                "total_capacity": {"cpu_cores": cpu, "memory_mib": BIG_MEM,
                                   "storage_gb": BIG_STO},
                "topology": {"ag": ag}, "network": network,
            })
    proc_types = [{"type_id": s, "capacity": {"cpu_cores": cpu,
                   "memory_mib": BIG_MEM, "storage_gb": BIG_STO}}
                  for s, cpu in sku_cpu.items()]
    proc_caps = [{"bucket": ag, "max_bm": mx, "network": network}
                 for (f, network, ag, mx) in caps if f == fab]
    return {
        "demand_book": db, "in_stock": in_stock,
        "procurement_types": proc_types, "procurement_caps": proc_caps,
        "config": {"procurement_spread_dimension": "ag",
                   "max_solve_time_seconds": 10, "target_spread": {"ag": 3}},
    }


def post(req):
    data = json.dumps(req).encode()
    r = urllib.request.Request(URL, data=data,
                               headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(r, timeout=60) as resp:
        return json.loads(resp.read())


for fab in fabs:
    req = build_request(fab)
    print(f"\n===== fab {fab}: {len(req['demand_book'])} demand rows, "
          f"{len(req['in_stock'])} in-stock BM, {len(req['procurement_caps'])} caps =====")
    try:
        rep = post(req)
    except urllib.error.HTTPError as e:
        print("HTTP", e.code, e.read().decode()[:500])
        continue
    print("success:", rep["success"])
    for pf in rep["by_fab_period"]:
        proc = ", ".join(f"{d['type_id']}×{d['count']}" for d in pf["procurement"]) or "—"
        sd = sum(d["count"] for d in pf["split_decisions"])
        bal = pf.get("balance_after", {})
        sf = "; ".join(s["message"][:60] for s in pf.get("shortfalls", []))
        print(f"  [{'OK ' if pf['success'] else '缺口'}] {pf['period']}: "
              f"node_adds={pf['node_adds_total']} split={sd}VM 採購=[{proc}] "
              f"in_stock_used={pf['in_stock_bm_used']} balance_ag={bal}"
              + (f" 缺口:{sf}" if sf else ""))
