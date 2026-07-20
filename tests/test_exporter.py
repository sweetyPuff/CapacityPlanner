from pathlib import Path

from openpyxl import load_workbook

from captool.exporter import export_summary, generate_v2_template
from captool.importer import import_legacy
from captool.planner import run_check
from captool.solver.naive import NaiveSolver

FIXTURE = Path(__file__).parent / "fixtures" / "sample_legacy.xlsx"


def test_export_summary(tmp_path):
    pi = import_legacy(FIXTURE)
    result = run_check(pi, NaiveSolver(), mode="conservative")
    out = tmp_path / "out.xlsx"
    export_summary(pi, result, out)
    wb = load_workbook(out)
    assert "summary" in wb.sheetnames and "placements" in wb.sheetnames
    ws = wb["summary"]
    headers = [c.value for c in ws[1]]
    assert "狀態" in headers and "缺口 vcore" in headers
    # 6 pools × 6 months = 36 資料列(+1 標頭 +1 附註)
    data_rows = [r for r in ws.iter_rows(min_row=2) if r[0].value in ("A", "B", "C")]
    assert len(data_rows) == 36


def test_export_summary_marks_gap(tmp_path):
    pi = import_legacy(FIXTURE)
    result = run_check(pi, NaiveSolver(), mode="conservative")
    out = tmp_path / "out.xlsx"
    export_summary(pi, result, out)
    ws = load_workbook(out)["summary"]
    statuses = {(r[0].value, r[1].value, r[2].value): r for r in ws.iter_rows(min_row=2)
                if r[0].value in ("A", "B", "C")}
    header = [c.value for c in ws[1]]
    status_idx = header.index("狀態")
    # B/network1 2026-07 在 Excel 中即為負庫存 → 工具也應標缺口
    assert statuses[("B", "network1", "2026-07")][status_idx].value == "缺口"


def test_generate_v2_template(tmp_path):
    pi = import_legacy(FIXTURE)
    out = tmp_path / "template_v2.xlsx"
    generate_v2_template(pi, out)
    wb = load_workbook(out)
    for sheet in ("HW_SKU", "HW_Current", "HW_MoveIn", "VM_Spec", "README", "A", "B", "C"):
        assert sheet in wb.sheetnames, sheet
    ws = wb["HW_SKU"]
    assert [c.value for c in ws[1]] == ["name", "vcore_per_node", "usable_ratio"]
    assert ws["A2"].value == "default-64"
    # 廠區 tab Return 區含機型欄
    wa = wb["A"]
    assert wa["M4"].value == "機型"
    assert wa["M5"].value == "default-64"
    # v2 範本可被 detect_format 認出
    from captool.importer import detect_format
    assert detect_format(out) == "v2"
