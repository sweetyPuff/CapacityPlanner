"""黃金測試:excel_compat 模式必須與現行 Excel 公式逐格一致(遷移正確性)。"""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from captool.importer import import_legacy
from captool.models import Pool
from captool.planner import run_check
from captool.solver.naive import NaiveSolver

FIXTURE = Path(__file__).parent / "fixtures" / "sample_legacy.xlsx"

# Excel 只算到 2026-09(AD 欄);AB=2026-07 起
MONTH_COLS = {"2026-07": "AB", "2026-08": "AC", "2026-09": "AD"}


def test_instock_matches_excel_cached_values():
    plan_input = import_legacy(FIXTURE)
    result = run_check(plan_input, NaiveSolver(), mode="excel_compat")
    wb = load_workbook(FIXTURE, data_only=True)
    ws = wb["summary"]
    checked = 0
    for row in range(3, 9):
        pool = Pool(fab=str(ws[f"A{row}"].value), bm_group=str(ws[f"B{row}"].value))
        for month, col in MONTH_COLS.items():
            expected = ws[f"{col}{row}"].value
            outcome = next(o for o in result.outcomes
                           if o.pool == pool and o.month == month)
            assert outcome.stock_float == pytest.approx(expected), (
                f"{pool.label} {month}: 工具 {outcome.stock_float} != Excel {expected}")
            checked += 1
    assert checked == 18


def test_reproduces_negative_stock_gap():
    plan_input = import_legacy(FIXTURE)
    result = run_check(plan_input, NaiveSolver(), mode="excel_compat")
    b1 = Pool(fab="B", bm_group="network1")
    o7 = next(o for o in result.for_pool(b1) if o.month == "2026-07")
    assert o7.stock_float == pytest.approx(-7.36328125)
    assert not o7.feasible
