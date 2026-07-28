import math
from pathlib import Path

from openpyxl import Workbook, load_workbook

from captool.exporter import generate_v2_template
from captool.importer import import_any, import_legacy, import_v2
from captool.models import (NodeReturn, PlanInput, Pool, ProductPolicy, Sku,
                            VmSpecDemand)
from captool.planner import run_check
from captool.solver.naive import NaiveSolver

FIXTURE = Path(__file__).parent / "fixtures" / "sample_legacy.xlsx"


def _v2_file(tmp_path):
    path = tmp_path / "v2.xlsx"
    generate_v2_template(import_legacy(FIXTURE), path)
    return path


def test_roundtrip_coarse_demand(tmp_path):
    path = _v2_file(tmp_path)
    original = import_legacy(FIXTURE)
    pi = import_v2(path)
    a1 = Pool(fab="A", bm_group="network1")
    assert pi.demand_vcore(a1, "2026-07") == original.demand_vcore(a1, "2026-07")
    assert pi.return_by_sku(a1, "2026-08") == original.return_by_sku(a1, "2026-08")
    assert pi.months == original.months
    # 示範列不得成為資料
    assert all("示範" not in d.product for d in pi.demands)
    assert pi.vm_demands == []            # 示範 detail 列被跳過


def test_detail_product_derives_count_and_policy(tmp_path):
    path = _v2_file(tmp_path)
    wb = load_workbook(path)
    wa = wb["A"]
    # 找一個空列加入 detail product:VM vcore=60, 每台上限=2, 共居=teamB, 2026-07=180
    r = wa.max_row + 1
    wa.cell(row=r, column=1, value="db")
    wa.cell(row=r, column=2, value="network1")
    wa.cell(row=r, column=3, value=60)
    wa.cell(row=r, column=4, value=2)
    wa.cell(row=r, column=5, value="teamB")
    wa.cell(row=r, column=6, value=180)     # 180/60 = 3 台
    wb.save(path)
    pi = import_v2(path)
    a1 = Pool(fab="A", bm_group="network1")
    assert pi.vm_batch(a1, "2026-07") == [(60, 3)]
    pol = pi.policy_for(a1, "db")
    assert pol is not None
    assert pol.vm_size_vcore == 60 and pol.max_per_machine == 2
    assert pol.co_residency == "teamB"


def test_detail_non_multiple_warns_and_rounds_up(tmp_path):
    path = _v2_file(tmp_path)
    wb = load_workbook(path)
    wa = wb["A"]
    r = wa.max_row + 1
    wa.cell(row=r, column=1, value="web")
    wa.cell(row=r, column=2, value="network1")
    wa.cell(row=r, column=3, value=60)
    wa.cell(row=r, column=6, value=100)     # 100/60 → 進位 2 台
    wb.save(path)
    pi = import_v2(path)
    a1 = Pool(fab="A", bm_group="network1")
    assert pi.vm_batch(a1, "2026-07") == [(60, 2)]
    assert any("整數倍" in i.message and i.severity == "warning" for i in pi.issues)


def test_exclusive_and_free_policy(tmp_path):
    path = _v2_file(tmp_path)
    wb = load_workbook(path)
    wa = wb["A"]
    r = wa.max_row + 1
    wa.cell(row=r, column=1, value="iso"); wa.cell(row=r, column=2, value="network1")
    wa.cell(row=r, column=3, value=51); wa.cell(row=r, column=5, value="獨佔")
    wa.cell(row=r, column=6, value=51)
    wb.save(path)
    pi = import_v2(path)
    pol = pi.policy_for(Pool(fab="A", bm_group="network1"), "iso")
    assert pol.co_residency == "exclusive"
    assert pol.max_per_machine is None      # D 空 = None


def test_import_any_dispatches_new(tmp_path):
    path = _v2_file(tmp_path)
    assert import_any(path).months == import_v2(path).months


def test_menu_column_worker_and_newbuild(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "menu.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    a = wb.create_sheet("A")
    a["A4"], a["B4"], a["C4"], a["F4"] = "Product", "BM Group", "VM vcore", "Menu"
    a["G4"], a["H4"] = "2026-07", "2026-08"      # Menu 欄在 F,月份右移到 G
    # worker 列
    a["A5"], a["B5"], a["F5"], a["G5"] = "web", "network1", "Worker", 200
    # new build 列(Menu=B,月份=cluster 數)
    a["A6"], a["B6"], a["F6"], a["H6"] = "c-x", "network1", "B", 1
    hs = wb.create_sheet("HW_SKU")
    hs.append(["name", "vcore_per_node", "usable_ratio"])
    hs.append(["std-64", 64, 0.8])
    wb.create_sheet("HW_Current").append(["fab", "bm_group", "sku", "count"])
    wb.create_sheet("HW_MoveIn").append(
        ["fab", "bm_group", "sku", "month", "count"])
    cm = wb.create_sheet("Cluster_Menu")
    cm.append(["menu", "role", "count", "co_residency", "vm_vcore"])
    cm.append(["B", "master", 5, "shared", 16])
    wb.save(path)

    pi = import_v2(path)
    a1 = Pool(fab="A", bm_group="network1")
    # worker 列 → 需求;month 從 G 讀對
    assert pi.demand_vcore(a1, "2026-07") == 200
    # new build 列 → 不進 demands,進 new_builds
    assert all(d.product != "c-x" for d in pi.demands)
    assert len(pi.new_builds) == 1
    nb = pi.new_builds[0]
    assert (nb.cluster, nb.menu, nb.month, nb.count) == ("c-x", "B", "2026-08", 1)
    assert pi.menus["B"] == [("master", 5, "shared", 16)]


def test_hw_current_ag_and_caps(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "ag.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    a = wb.create_sheet("A")
    a["A4"], a["B4"], a["C4"] = "Product", "BM Group", "VM vcore"
    a["F4"] = "2026-07"
    a["A5"], a["B5"], a["F5"] = "web", "network1", 100
    hs = wb.create_sheet("HW_SKU")
    hs.append(["name", "vcore_per_node", "usable_ratio"])
    hs.append(["std-64", 64, 0.8])
    hc = wb.create_sheet("HW_Current")
    hc.append(["fab", "bm_group", "sku", "count", "ag"])
    hc.append(["A", "network1", "std-64", 2, "ag1"])
    hc.append(["A", "network1", "std-64", 3, "ag2"])
    wb.create_sheet("HW_MoveIn").append(
        ["fab", "bm_group", "sku", "month", "count"])
    hcap = wb.create_sheet("HW_Caps")
    hcap.append(["fab", "network", "ag", "max_bm"])
    hcap.append(["A", "network1", "ag1", 20])
    hcap.append(["A", "network1", "ag2", 15])
    wb.save(path)

    pi = import_v2(path)
    assert sorted({c.ag for c in pi.currents}) == ["ag1", "ag2"]
    assert sum(c.count for c in pi.currents) == 5
    a1 = Pool(fab="A", bm_group="network1")
    caps = {cap.ag: cap.max_bm for cap in pi.caps if cap.pool == a1}
    assert caps == {"ag1": 20, "ag2": 15}


def test_hw_current_without_ag_defaults_empty(tmp_path):
    # 舊檔 HW_Current 無 ag 欄 → ag 預設 "",caps 為空
    path = _v2_file(tmp_path)   # generate_v2_template 產的 4 欄 HW_Current
    pi = import_v2(path)
    assert all(c.ag == "" for c in pi.currents)
    assert pi.caps == []


def test_old_format_still_imports(tmp_path):
    # 手工建一個含 VM_Spec 分頁的舊 v2,確認相容路徑仍可解析
    path = tmp_path / "old_v2.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    a = wb.create_sheet("A")
    a["C3"] = "User Demand (vcore)"
    a["A4"], a["B4"], a["C4"] = "Product", "BM Group", "2026-07"
    a["A5"], a["B5"], a["C5"] = "svc", "network1", 100
    a["K4"], a["L4"], a["M4"], a["N4"] = "Product", "BM Group", "機型", "2026-07"
    hs = wb.create_sheet("HW_SKU"); hs.append(["name", "vcore_per_node", "usable_ratio"]); hs.append(["std-64", 64, 0.8])
    hc = wb.create_sheet("HW_Current"); hc.append(["fab", "bm_group", "sku", "count"]); hc.append(["A", "network1", "std-64", 5])
    hm = wb.create_sheet("HW_MoveIn"); hm.append(["fab", "bm_group", "sku", "month", "count"])
    vs = wb.create_sheet("VM_Spec"); vs.append(["fab", "bm_group", "product", "month", "vm_size_vcore", "count"]); vs.append(["A", "network1", "ai", "2026-07", 32, 2])
    wb.save(path)
    pi = import_v2(path)
    a1 = Pool(fab="A", bm_group="network1")
    assert pi.demand_vcore(a1, "2026-07") == 100
    assert pi.vm_batch(a1, "2026-07") == [(32, 2)]   # 舊 VM_Spec 分頁仍解析


def test_vm_spec_only_pool_is_planned(tmp_path):
    """I-1: 只出現在 VM_Spec 的 pool 也必須進入 pi.pools 並被排入推演。"""
    path = _v2_file(tmp_path)
    wb = load_workbook(path)
    # 新格式不產生 VM_Spec,需手動建立以測試向後相容性
    ws = wb.create_sheet("VM_Spec")
    ws.append(["fab", "bm_group", "product", "month", "vm_size_vcore", "count"])
    new_pool = Pool(fab="A", bm_group="network9")
    ws.append(["A", "network9", "Product Only-In-VmSpec", "2026-07", 32, 4])
    wb.save(path)
    pi = import_v2(path)
    assert new_pool in pi.pools
    result = run_check(pi, NaiveSolver(), mode="conservative")
    outcomes = result.for_pool(new_pool)
    assert len(outcomes) == len(pi.months)
    assert any(o.vm_demand == [(32, 4)] for o in outcomes)


def test_non_numeric_count_reported_not_raised(tmp_path):
    """I-5: HW_MoveIn 的 count 非數值時應回報錯誤並以 0 計,而非整個 raise 中斷匯入。"""
    path = _v2_file(tmp_path)
    wb = load_workbook(path)
    wb["HW_MoveIn"].append(["A", "network1", "default-64", "2026-09", "abc"])
    wb.save(path)
    pi = import_v2(path)  # 不應 raise
    errors = [i for i in pi.issues if i.severity == "error" and i.sheet == "HW_MoveIn"]
    assert any("應為數值" in i.message and i.cell for i in errors)
    # 非數值列以 0 計,既有 2026-09 default-64 movein(來自 summary Q3=25)不受影響
    assert pi.movein_by_sku(Pool(fab="A", bm_group="network1"), "2026-09") == {"default-64": 25}


def test_multi_sku_returns_roundtrip(tmp_path):
    """I-3: 同一 product/bm_group 但不同機型的 Return 應各自成列,往返後不互相覆蓋。"""
    skus = {"s1": Sku(name="s1", vcore_per_node=64, usable_ratio=0.8),
            "s2": Sku(name="s2", vcore_per_node=128, usable_ratio=0.8)}
    pool = Pool(fab="A", bm_group="network1")
    pi = PlanInput(
        skus=skus, months=["2026-07"], pools=[pool],
        demands=[], vm_demands=[], moveins=[],
        returns=[
            NodeReturn(pool=pool, product="P1", sku_name="s1", month="2026-07", count=5),
            NodeReturn(pool=pool, product="P1", sku_name="s2", month="2026-07", count=7),
        ],
        currents=[])
    path = tmp_path / "v2.xlsx"
    generate_v2_template(pi, path)
    pi2 = import_v2(path)
    assert pi2.return_by_sku(pool, "2026-07") == {"s1": 5, "s2": 7}


def test_detail_rows_roundtrip_through_template(tmp_path):
    """generate_v2_template 必須自行重寫 VM-detail + policy 資料列(不靠手動補 VM_Spec)。"""
    skus = {"s1": Sku(name="s1", vcore_per_node=64, usable_ratio=0.8)}
    pool1 = Pool(fab="A", bm_group="network1")
    pool2 = Pool(fab="A", bm_group="network2")
    pi = PlanInput(
        skus=skus, months=["2026-07", "2026-08"], pools=[pool1, pool2],
        demands=[],
        vm_demands=[
            VmSpecDemand(pool=pool1, product="Product X", month="2026-07",
                        vm_size_vcore=32, count=4),
            VmSpecDemand(pool=pool2, product="Product Y", month="2026-08",
                        vm_size_vcore=16, count=10),
        ],
        moveins=[], returns=[], currents=[],
        policies=[
            ProductPolicy(pool=pool1, product="Product X", vm_size_vcore=32,
                         max_per_machine=2, co_residency="exclusive"),
            ProductPolicy(pool=pool2, product="Product Y", vm_size_vcore=16,
                         max_per_machine=None, co_residency="free"),
        ])
    path = tmp_path / "v2.xlsx"
    generate_v2_template(pi, path)
    assert "VM_Spec" not in load_workbook(path).sheetnames   # 無需獨立分頁
    pi2 = import_v2(path)
    assert pi2.vm_batch(pool1, "2026-07") == [(32, 4)]
    assert pi2.vm_batch(pool2, "2026-08") == [(16, 10)]
    pol1 = pi2.policy_for(pool1, "Product X")
    assert pol1 is not None
    assert pol1.vm_size_vcore == 32
    assert pol1.max_per_machine == 2
    assert pol1.co_residency == "exclusive"
    pol2 = pi2.policy_for(pool2, "Product Y")
    assert pol2 is not None
    assert pol2.vm_size_vcore == 16
    assert pol2.max_per_machine is None
    assert pol2.co_residency == "free"


def test_detail_bad_vm_size_reported_not_raised(tmp_path):
    """Guard C (VM vcore):非正整數時應回報錯誤並跳過該列,而非整個 raise 中斷匯入。"""
    path = _v2_file(tmp_path)
    wb = load_workbook(path)
    wa = wb["A"]
    r = wa.max_row + 1
    wa.cell(row=r, column=1, value="bad")           # product
    wa.cell(row=r, column=2, value="network1")      # group
    wa.cell(row=r, column=3, value="abc")           # C: non-numeric VM vcore
    wa.cell(row=r, column=4, value=2)               # D: max_per
    wa.cell(row=r, column=6, value=100)             # F: some month vcore
    wb.save(path)
    pi = import_v2(path)  # 不應 raise
    errors = [i for i in pi.issues if i.severity == "error" and i.sheet == "A"]
    assert any("VM vcore" in i.message and "非正整數" in i.message for i in errors)
    # 該列已被跳過,不應產生 vm_demands 或 policy
    assert not any(d.product == "bad" for d in pi.vm_demands)
    assert not any(p.product == "bad" for p in pi.policies)


def test_new_format_allzero_month_kept(tmp_path):
    """新格式月份欄若整欄皆為空白/0,仍應保留在 pi.months(不可被 vcore 判斷排除)。"""
    path = tmp_path / "new_v2.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    a = wb.create_sheet("A")
    a["A4"], a["B4"] = "Product", "BM Group"
    a["C4"], a["D4"], a["E4"] = "VM vcore", "每台上限", "共居"
    a["F4"], a["G4"] = "2026-07", "2026-08"    # 2026-08 整欄空白
    a["A5"], a["B5"] = "svc", "network1"
    a["F5"] = 100
    a.cell(row=4, column=20, value="Product")
    a.cell(row=4, column=21, value="BM Group")
    a.cell(row=4, column=22, value="機型")
    hs = wb.create_sheet("HW_SKU"); hs.append(["name", "vcore_per_node", "usable_ratio"]); hs.append(["std-64", 64, 0.8])
    hc = wb.create_sheet("HW_Current"); hc.append(["fab", "bm_group", "sku", "count"]); hc.append(["A", "network1", "std-64", 5])
    hm = wb.create_sheet("HW_MoveIn"); hm.append(["fab", "bm_group", "sku", "month", "count"])
    wb.save(path)
    pi = import_v2(path)
    assert "2026-08" in pi.months
    assert pi.demand_vcore(Pool(fab="A", bm_group="network1"), "2026-08") == 0
